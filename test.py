import os
from datetime import datetime
import pandas as pd
import streamlit as st

# --- SEITEN-STYLING (iOS Minimal-Look) ---
st.set_page_config(page_title="Problem Melden", page_icon="🚨", layout="centered")

st.markdown("""
    <style>
    .stApp { background-color: #f8f9fa; }
    .app-title { color: #1E67C6; text-align: center; font-size: 28px; font-weight: bold; margin-bottom: 25px; }
    div[data-baseweb="textarea"] {
        border-radius: 15px !important;
        border: 1px solid #E0E0E0 !important;
        box-shadow: 0px 2px 10px rgba(0, 0, 0, 0.05) !important;
        background-color: white !important;
    }
    div.stButton > button {
        background-color: #007AFF !important; color: white !important;
        border-radius: 10px !important; border: none !important;
        width: 100% !important; padding: 12px !important;
        font-size: 18px !important; font-weight: 500 !important;
        box-shadow: 0px 4px 6px rgba(0, 122, 255, 0.2) !important;
    }
    #MainMenu, footer, header { visibility: hidden; }
    </style>
""", unsafe_allow_html=True)

st.markdown('<p class="app-title">Problem Melden</p>', unsafe_allow_html=True)

# --- SITZPLATZ-PRÜFUNG ---
query_params = st.query_params
qr_platz = query_params.get("platz", None)

if qr_platz is None:
    st.error("❌ Zugriff verweigert: Kein Sitzplatz erkannt.")
    st.warning("Bitte nutzen Sie die offizielle Desktop-Verknüpfung oder scannen Sie den QR-Code an Ihrem Platz.")
    st.stop()

st.info(f"📍 **Sitzplatz: {qr_platz}**")
st.write("")

problem_text = st.text_area(label="Problembeschreibung", placeholder="Beschreibe dein Problem hier...", label_visibility="collapsed", height=200)

# --- AUTOMATISCHER SPEICHERPFAD (Direkt in Ihr Firmen-OneDrive) ---
user_profile = os.environ.get("USERPROFILE")
if user_profile:
    ONEDRIVE_ORDNER = os.path.join(user_profile, r"OneDrive - IBB Institut für Berufliche Bildung AG - Viona")
    EXCEL_DATEI = os.path.join(ONEDRIVE_ORDNER, "ProblemMeldungen_Sync.xlsx")
else:
    st.error("Fehler: Windows-Benutzerprofil (USERPROFILE) konnte nicht ermittelt werden.")
    st.stop()

# WICHTIG: Stelle sicher, dass diese Imports ganz oben oder hier stehen:
from openpyxl.styles import PatternFill, Border, Side, Font, Alignment
from openpyxl.utils import get_column_letter

# --- ABSENDEN BUTTON ---
if st.button("Senden"):
    if problem_text.strip() == "":
        st.warning("Bitte beschreibe zuerst das Problem.")
    else:
        zeitpunkt = datetime.now().strftime("%d.%m.%Y %H:%M:%S")
        neue_daten = pd.DataFrame([{"Zeitstempel": zeitpunkt, "Sitzplatz": str(qr_platz), "Problem": str(problem_text)}])
       
        try:
            # Falls der OneDrive-Ordner gar nicht existiert, erstellen wir ihn
            if not os.path.exists(ONEDRIVE_ORDNER):
                os.makedirs(ONEDRIVE_ORDNER)

            # Prüfen, ob die Datei existiert UND nicht komplett leer ist
            if os.path.exists(EXCEL_DATEI) and os.path.getsize(EXCEL_DATEI) > 0:
                try:
                    alt_daten = pd.read_excel(EXCEL_DATEI, skiprows=1, usecols="B:D")
                    kombiniert = pd.concat([alt_daten, neue_daten], ignore_index=True)
                except Exception:
                    kombiniert = neue_daten
            else:
                kombiniert = neue_daten
               
            # 1. Daten in Excel schreiben (startet bei B2)
            with pd.ExcelWriter(EXCEL_DATEI, engine="openpyxl") as writer:
                kombiniert.to_excel(writer, index=False, startrow=1, startcol=1, sheet_name="Tabelle1")
                
                # 2. AUTOMATISCHES STYLING
                sheet = writer.book.active
                
                # Definition der Styles
                header_fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
                white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
                blue_fill = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
                
                thin_side = Side(border_style="thin", color="000000")
                cell_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
                
                center_alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)
                
                # Formatierung der Zeilen und Zellen (Schriftgröße 14)
                for row_idx, row in enumerate(sheet.iter_rows(min_row=2, max_row=sheet.max_row, min_col=2, max_col=4), start=2):
                    for cell in row:
                        cell.border = cell_border
                        cell.alignment = center_alignment
                        
                        if row_idx == 2:
                            cell.fill = header_fill
                            cell.font = Font(name="Calibri", size=14, bold=True)
                        else:
                            cell.font = Font(name="Calibri", size=14, bold=False)
                            if row_idx % 2 == 0:
                                cell.fill = blue_fill
                            else:
                                cell.fill = white_fill

                # 3. AUTOMATISCHE SPALTENBREITE ANPASSEN (Individuell für jede Spalte)
                for col in sheet.iter_cols(min_col=2, max_col=4):
                    col_letter = get_column_letter(col[0].column)
                    max_len = 0
                    
                    if col_letter == "C":
                        header_value = sheet["C2"].value
                        max_len = len(str(header_value)) if header_value else 9
                    else:
                        for cell in col[1:]:
                            if cell.value:
                                lines = str(cell.value).split('\n')
                                for line in lines:
                                    if len(line) > max_len:
                                        max_len = len(line)
                    
                    sheet.column_dimensions[col_letter].width = max(max_len + 4, 12)
                                
            st.success(f"🎉 Vielen Dank! Das Problem für {qr_platz} wurde erfolgreich übermittelt und perfekt formatiert.")
            
        except Exception as e:
            st.error(f"Fehler beim Schreiben in OneDrive: {e}")
            st.info("Bitte prüfen Sie, ob die Excel-Datei gerade geöffnet ist. Schließen Sie sie und versuchen Sie es noch einmal.")